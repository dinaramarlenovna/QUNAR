from __future__ import annotations

from io import BytesIO
from datetime import datetime
import html
import re
import pandas as pd


def _clean(v):
    if v is None:
        return ""
    return str(v)


def _safe_sheet_name(name: str, fallback: str = "Report") -> str:
    name = re.sub(r"[\\/*?:\[\]]", " ", str(name or fallback)).strip()
    return (name[:31] or fallback)


def _num(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value or "").replace("\u00a0", " ").replace(",", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def _unique_inputs(payloads):
    seen = set()
    out = []
    for payload in payloads:
        for row in payload.get("inputs", []) or []:
            if len(row) < 3:
                continue
            key = (str(row[0]), str(row[1]), str(row[2]))
            if key not in seen:
                seen.add(key)
                out.append(list(row[:3]))
    return out


def _consolidated_input_groups(payloads):
    """One shared assumption table plus explicit differences grouped by report.

    Fields that vary across report sections (such as a 10 ha field and 100 ha
    optimization) are never collapsed into an ambiguous common value.
    """
    from collections import OrderedDict
    per_mode = []
    for payload in payloads:
        rows = OrderedDict()
        for row in payload.get('inputs', []) or []:
            if len(row) >= 3:
                rows[str(row[0])] = (str(row[1]), str(row[2]))
        per_mode.append(rows)
    if not per_mode:
        return [], []
    counts = {}
    for rows in per_mode:
        for name, pair in rows.items():
            counts.setdefault(name, []).append(pair)
    common = []
    for name, pairs in counts.items():
        if len(pairs) == len(per_mode) and len(set(pairs)) == 1:
            common.append((name, *pairs[0]))
    # Also consolidate identical values appearing in 2+ sections, even when
    # another section does not define that parameter.
    for name, pairs in counts.items():
        if name not in {r[0] for r in common} and len(pairs) >= 2 and len(set(pairs)) == 1:
            common.append((name, *pairs[0]))
    common_names = {r[0] for r in common}
    specifics = []
    for index, (payload, rows) in enumerate(zip(payloads, per_mode), 1):
        unique = [(name, *pair) for name, pair in rows.items() if name not in common_names]
        if unique:
            specifics.append((f"{index}. {payload.get('mode_label','Mode')}", unique))
    return common, specifics


def _xlsx_formats(wb):
    return {
        "brand": wb.add_format({"bold": True, "font_size": 24, "font_color": "#1F5135"}),
        "subtitle": wb.add_format({"bold": True, "font_size": 11, "font_color": "#52665A"}),
        "section": wb.add_format({"bold": True, "font_size": 16, "font_color": "#20372A", "bottom": 2, "bottom_color": "#8BAD96"}),
        "header": wb.add_format({"bold": True, "bg_color": "#EAF3EC", "font_color": "#234F36", "border": 1, "border_color": "#D7E3DA", "text_wrap": True, "valign": "vcenter"}),
        "cell": wb.add_format({"valign": "top", "text_wrap": True, "border": 1, "border_color": "#E5EAE6"}),
        "key": wb.add_format({"bold": True, "bg_color": "#F1F7F2", "font_color": "#234F36", "text_wrap": True, "border": 1, "border_color": "#D7E3DA", "valign": "top"}),
        "recommend": wb.add_format({"bold": True, "bg_color": "#EEF7F0", "font_color": "#183A25", "text_wrap": True, "border": 1, "border_color": "#BFD5C5", "valign": "top"}),
        "small": wb.add_format({"font_size": 9, "font_color": "#66736B", "text_wrap": True}),
        "money": wb.add_format({"num_format": '#,##0', "border": 1, "border_color": "#E5EAE6"}),
        "pct": wb.add_format({"num_format": '0.0', "border": 1, "border_color": "#E5EAE6"}),
    }


def _write_payload_sheet(writer, sheet_name: str, payload: dict, include_inputs: bool = True):
    wb = writer.book
    fm = _xlsx_formats(wb)
    ws = wb.add_worksheet(_safe_sheet_name(sheet_name))
    writer.sheets[ws.name] = ws
    ws.hide_gridlines(2)
    ws.freeze_panes(5, 0)
    ws.set_column("A:A", 30)
    ws.set_column("B:B", 28)
    ws.set_column("C:C", 58)
    ws.set_row(0, 34)

    ws.write("A1", "QŪNAR", fm["brand"])
    ws.merge_range("B1:C1", payload.get("mode_label", ""), fm["section"])
    ws.merge_range("A2:C2", payload.get("brand_translation", ""), fm["subtitle"])
    ws.write("A3", payload.get("generated_label", "Generated"), fm["small"])
    ws.write("B3", payload.get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M")), fm["small"])

    row = 5
    summary = payload.get("summary", []) or []
    if summary:
        ws.merge_range(row, 0, row, 2, payload.get("summary_title", "Key results"), fm["section"]); row += 2
        ws.write_row(row, 0, [payload.get("indicator_label", "Indicator"), payload.get("value_label", "Value"), payload.get("note_label", "Note")], fm["header"]); row += 1
        for item in summary:
            ws.write(row, 0, _clean(item.get("label")), fm["key"])
            value = item.get("value", "")
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                ws.write_number(row, 1, float(value), fm["money"] if item.get("kind") == "money" else fm["cell"])
            else:
                ws.write(row, 1, _clean(value), fm["cell"])
            ws.write(row, 2, _clean(item.get("note")), fm["cell"])
            row += 1

    recommendations = payload.get("recommendations", []) or []
    if recommendations:
        row += 1
        ws.merge_range(row, 0, row, 2, payload.get("recommendations_title", "Recommendations"), fm["section"]); row += 2
        for i, rec in enumerate(recommendations, 1):
            if isinstance(rec, dict):
                title = rec.get("title") or f"{i}."
                text = rec.get("text") or rec.get("recommendation") or ""
                evidence = rec.get("evidence") or rec.get("why") or ""
                effect = rec.get("effect") or ""
                ws.write(row, 0, f"{i}. {title}", fm["recommend"])
                ws.write(row, 1, _clean(text), fm["recommend"])
                ws.write(row, 2, "\n".join(x for x in [_clean(evidence), _clean(effect)] if x), fm["cell"])
            else:
                ws.write(row, 0, f"{i}.", fm["recommend"])
                ws.merge_range(row, 1, row, 2, _clean(rec), fm["recommend"])
            row += 1

    analysis = payload.get("analysis", []) or []
    if analysis:
        row += 1
        ws.merge_range(row, 0, row, 2, payload.get("analysis_title", "Analysis"), fm["section"]); row += 2
        for text in analysis:
            ws.merge_range(row, 0, row, 2, _clean(text), fm["cell"])
            row += 1

    table = payload.get("table")
    if table:
        row += 1
        ws.merge_range(row, 0, row, 2, payload.get("table_title", "Detailed results"), fm["section"]); row += 2
        table_df = pd.DataFrame(table)
        if not table_df.empty:
            start = row
            table_df.to_excel(writer, sheet_name=ws.name, startrow=start, startcol=0, index=False)
            for c, col in enumerate(table_df.columns):
                ws.write(start, c, col, fm["header"])
                ws.set_column(c, c, min(max(14, len(str(col)) + 2), 25))
            row += len(table_df) + 2

    # Analytical chart in Excel: only where a meaningful numeric comparison exists.
    chart = payload.get("chart") or {}
    if chart.get("categories") and chart.get("values"):
        row += 1
        chart_start = row
        ws.write_row(chart_start, 0, [chart.get("category_label", "Category"), chart.get("value_label", "Value")], fm["header"])
        for i, (cat, val) in enumerate(zip(chart["categories"], chart["values"]), 1):
            ws.write(chart_start + i, 0, _clean(cat), fm["cell"])
            ws.write_number(chart_start + i, 1, float(val), fm["pct"])
        ch = wb.add_chart({"type": "bar"})
        ch.add_series({
            "name": chart.get("title", "QŪNAR"),
            "categories": [ws.name, chart_start + 1, 0, chart_start + len(chart["categories"]), 0],
            "values": [ws.name, chart_start + 1, 1, chart_start + len(chart["values"]), 1],
            "data_labels": {"value": True},
        })
        ch.set_title({"name": chart.get("title", "QŪNAR analytical view")})
        ch.set_legend({"none": True})
        ch.set_style(10)
        ch.set_size({"width": 720, "height": 380})
        ws.insert_chart(chart_start, 3, ch)
        row += max(len(chart["categories"]) + 2, 20)

    if include_inputs and payload.get("inputs"):
        row += 1
        ws.merge_range(row, 0, row, 2, payload.get("inputs_title", "Inputs and sources"), fm["section"]); row += 2
        ws.write_row(row, 0, [payload.get("parameter_label", "Parameter"), payload.get("value_label", "Value"), payload.get("source_label", "Source")], fm["header"]); row += 1
        for p, v, s in payload["inputs"]:
            ws.write(row, 0, _clean(p), fm["cell"])
            ws.write(row, 1, _clean(v), fm["cell"])
            ws.write(row, 2, _clean(s), fm["cell"])
            row += 1
    return ws


def build_excel_report(payload: dict) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        _write_payload_sheet(writer, "QUNAR_Report", payload, include_inputs=True)
    output.seek(0)
    return output.getvalue()


def build_combined_excel_report(payloads: list[dict]) -> bytes:
    payloads = [p for p in payloads if p]
    output = BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        wb = writer.book
        fm = _xlsx_formats(wb)
        summary_ws = wb.add_worksheet("Summary")
        writer.sheets["Summary"] = summary_ws
        summary_ws.hide_gridlines(2)
        summary_ws.set_column("A:A", 31)
        summary_ws.set_column("B:B", 54)
        summary_ws.set_column("C:C", 28)
        lang = (payloads[0].get("lang") if payloads else "en") or "en"
        combined_title = {"kk":"Біріктірілген аналитикалық есеп","ru":"Сводный аналитический отчет","en":"Integrated analytical report"}.get(lang,"Integrated analytical report")
        included_title = {"kk":"Қосылған бөлімдер","ru":"Включенные разделы","en":"Included sections"}.get(lang,"Included sections")
        summary_ws.write("A1", "QŪNAR", fm["brand"])
        summary_ws.merge_range("B1:C1", combined_title, fm["section"])
        summary_ws.write("A3", included_title, fm["section"])
        row = 5
        for i, payload in enumerate(payloads, 1):
            summary_ws.write(row, 0, f"{i}. {payload.get('mode_label','')}", fm["key"])
            recs = payload.get("recommendations", []) or []
            first = recs[0] if recs else ""
            if isinstance(first, dict):
                first = first.get("text") or first.get("recommendation") or first.get("title") or ""
            summary_ws.write(row, 1, _clean(first), fm["recommend"])
            summary_ws.write(row, 2, payload.get("generated_at", ""), fm["cell"])
            row += 1

        for i, payload in enumerate(payloads, 1):
            _write_payload_sheet(writer, f"{i}_{payload.get('mode_label','Mode')}", payload, include_inputs=False)

        common, specifics = _consolidated_input_groups(payloads)
        if common or specifics:
            src=wb.add_worksheet('Sources');writer.sheets['Sources']=src
            src.hide_gridlines(2);src.set_column('A:A',35);src.set_column('B:B',30);src.set_column('C:C',72)
            rr=0
            def add_section(title,rows):
                nonlocal rr
                src.merge_range(rr,0,rr,2,title,fm['section']);rr+=1
                src.write_row(rr,0,['Parameter','Value','Source / role'],fm['header']);rr+=1
                for name,value,source in rows:
                    src.write_row(rr,0,[name,value,source],fm['cell']);rr+=1
                rr+=2
            if common:add_section('Shared parameters',common)
            for label,rows in specifics:add_section('Mode-specific: '+label,rows)
    output.seek(0)
    return output.getvalue()


def _register_pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        from matplotlib import font_manager
        regular = font_manager.findfont("DejaVu Sans")
        bold = font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight="bold"))
        pdfmetrics.registerFont(TTFont("QunarSans", regular))
        pdfmetrics.registerFont(TTFont("QunarSansBold", bold))
        return "QunarSans", "QunarSansBold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def _pdf_styles(regular, bold):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("QunarTitle", parent=styles["Title"], fontName=bold, fontSize=25, leading=29, textColor=colors.HexColor("#1F5135"), alignment=TA_CENTER, spaceAfter=8),
        "subtitle": ParagraphStyle("QunarSubtitle", parent=styles["BodyText"], fontName=bold, fontSize=10.5, leading=15, textColor=colors.HexColor("#53675A"), alignment=TA_CENTER, spaceAfter=15),
        "h2": ParagraphStyle("QunarH2", parent=styles["Heading2"], fontName=bold, fontSize=15, leading=19, textColor=colors.HexColor("#24342A"), spaceBefore=13, spaceAfter=8),
        "h3": ParagraphStyle("QunarH3", parent=styles["Heading3"], fontName=bold, fontSize=11.5, leading=15, textColor=colors.HexColor("#234F36"), spaceBefore=7, spaceAfter=4),
        "body": ParagraphStyle("QunarBody", parent=styles["BodyText"], fontName=regular, fontSize=10.5, leading=16, firstLineIndent=12, textColor=colors.HexColor("#202722"), spaceAfter=8),
        "key": ParagraphStyle("QunarKey", parent=styles["BodyText"], fontName=bold, fontSize=10.6, leading=15, textColor=colors.HexColor("#183A25"), backColor=colors.HexColor("#EEF7F0"), borderColor=colors.HexColor("#BFD5C5"), borderWidth=.7, borderPadding=8, spaceAfter=8),
        "small": ParagraphStyle("QunarSmall", parent=styles["BodyText"], fontName=regular, fontSize=8.4, leading=11.5, textColor=colors.HexColor("#67736B")),
    }


def _chart_image(chart: dict):
    if not chart or not chart.get("categories") or not chart.get("values"):
        return None
    import matplotlib.pyplot as plt
    cats = [str(x) for x in chart["categories"]]
    vals = [float(x) for x in chart["values"]]
    # Keep charts analytical and readable, not decorative.
    fig_h = max(2.6, min(5.2, 0.42 * len(cats) + 1.3))
    fig, ax = plt.subplots(figsize=(7.4, fig_h))
    y = list(range(len(cats)))
    ax.barh(y, vals)
    ax.set_yticks(y, labels=cats)
    ax.invert_yaxis()
    ax.set_title(chart.get("title", "QŪNAR analytical view"))
    ax.set_xlabel(chart.get("axis_label", chart.get("value_label", "")))
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", alpha=.18)
    for yi, v in zip(y, vals):
        ax.text(v, yi, f" {v:.1f}", va="center", fontsize=9)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _append_payload_pdf(story, payload: dict, styles, include_inputs: bool = True, include_brand: bool = True):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    body, h2, h3, key, small = styles["body"], styles["h2"], styles["h3"], styles["key"], styles["small"]
    if include_brand:
        story.append(Paragraph("QŪNAR", styles["title"]))
        story.append(Paragraph(html.escape(payload.get("brand_translation", "")), styles["subtitle"]))
    story.append(Paragraph(html.escape(payload.get("mode_label", "")), h2))
    story.append(Paragraph(html.escape(f"{payload.get('generated_label','Generated')}: {payload.get('generated_at','')}"), small))

    summary = payload.get("summary", []) or []
    if summary:
        story.append(Paragraph(html.escape(payload.get("summary_title", "Key results")), h2))
        # 2-column KPI table, compact and readable.
        data = []
        for x in summary:
            label = Paragraph(f"<b>{html.escape(str(x.get('label','')))}</b>", body)
            value = Paragraph(f"<b>{html.escape(str(x.get('value','')))}</b>", body)
            data.append([label, value])
        tbl = Table(data, colWidths=[76*mm, 84*mm], hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("VALIGN",(0,0),(-1,-1),"TOP"), ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#DCE5DE")),
            ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F6F9F7")), ("LEFTPADDING",(0,0),(-1,-1),7),
            ("RIGHTPADDING",(0,0),(-1,-1),7), ("TOPPADDING",(0,0),(-1,-1),7), ("BOTTOMPADDING",(0,0),(-1,-1),7)
        ]))
        story += [tbl, Spacer(1, 4*mm)]

    chart_buf = _chart_image(payload.get("chart") or {})
    if chart_buf:
        story.append(Paragraph(html.escape((payload.get("chart") or {}).get("section_title", "Analytical visualization")), h2))
        from reportlab.lib.utils import ImageReader
        iw, ih = ImageReader(chart_buf).getSize()
        chart_buf.seek(0)
        w = 166*mm
        h = w * (ih / iw)
        story.append(Image(chart_buf, width=w, height=h))
        story.append(Spacer(1, 2*mm))

    recommendations = payload.get("recommendations", []) or []
    if recommendations:
        story.append(Paragraph(html.escape(payload.get("recommendations_title", "Recommendations")), h2))
        for i, rec in enumerate(recommendations, 1):
            if isinstance(rec, dict):
                title = rec.get("title") or f"Recommendation {i}"
                text = rec.get("text") or rec.get("recommendation") or ""
                evidence = rec.get("evidence") or rec.get("why") or ""
                effect = rec.get("effect") or ""
                story.append(Paragraph(f"{i}. {html.escape(str(title))}", h3))
                story.append(Paragraph(html.escape(str(text)), key))
                if evidence:
                    story.append(Paragraph(f"<b>{html.escape(str(rec.get('evidence_label','Evidence')))}:</b> {html.escape(str(evidence))}", body))
                if effect:
                    story.append(Paragraph(f"<b>{html.escape(str(rec.get('effect_label','Expected effect')))}:</b> {html.escape(str(effect))}", body))
            else:
                story.append(Paragraph(f"{i}. {html.escape(str(rec))}", key))

    analysis = payload.get("analysis", []) or []
    if analysis:
        story.append(Paragraph(html.escape(payload.get("analysis_title", "Analysis")), h2))
        for text in analysis:
            story.append(Paragraph(html.escape(str(text)), body))

    table = payload.get("table")
    if table:
        story.append(Paragraph(html.escape(payload.get("table_title", "Detailed results")), h2))
        df = pd.DataFrame(table)
        if not df.empty:
            df = df.iloc[:, :8]
            data = [[Paragraph(f"<b>{html.escape(str(c))}</b>", small) for c in df.columns]]
            for _, row in df.head(15).iterrows():
                data.append([Paragraph(html.escape(_clean(v)), small) for v in row.tolist()])
            widths = [170*mm/len(df.columns)] * len(df.columns)
            tbl = Table(data, colWidths=widths, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EAF3EC")), ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D7E0D9")),
                ("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),3), ("RIGHTPADDING",(0,0),(-1,-1),3),
                ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4)
            ]))
            story.append(tbl)

    if include_inputs and payload.get("inputs"):
        story += [PageBreak(), Paragraph(html.escape(payload.get("inputs_title", "Inputs and sources")), h2)]
        data = [[Paragraph(f"<b>{html.escape(payload.get('parameter_label','Parameter'))}</b>", small), Paragraph(f"<b>{html.escape(payload.get('value_label','Value'))}</b>", small), Paragraph(f"<b>{html.escape(payload.get('source_label','Source'))}</b>", small)]]
        for p,v,s in payload["inputs"]:
            data.append([Paragraph(html.escape(_clean(p)), small), Paragraph(html.escape(_clean(v)), small), Paragraph(html.escape(_clean(s)), small)])
        tbl = Table(data, colWidths=[52*mm, 42*mm, 76*mm], repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EAF3EC")), ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D7E0D9")),
            ("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),4), ("RIGHTPADDING",(0,0),(-1,-1),4),
            ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4)
        ]))
        story.append(tbl)


def build_pdf_report(payload: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate
    regular, bold = _register_pdf_fonts()
    styles = _pdf_styles(regular, bold)
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm, title="QŪNAR report", author="QŪNAR")
    story = []
    _append_payload_pdf(story, payload, styles, include_inputs=True, include_brand=True)
    doc.build(story)
    out.seek(0)
    return out.getvalue()


def build_combined_pdf_report(payloads: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
    payloads = [p for p in payloads if p]
    regular, bold = _register_pdf_fonts()
    styles = _pdf_styles(regular, bold)
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm, title="QŪNAR integrated report", author="QŪNAR")
    lang = (payloads[0].get("lang") if payloads else "en") or "en"
    combined_title = {"kk":"Біріктірілген аналитикалық есеп","ru":"Сводный аналитический отчет","en":"Integrated analytical report"}.get(lang,"Integrated analytical report")
    structure_title = {"kk":"Есеп құрылымы","ru":"Структура отчета","en":"Report structure"}.get(lang,"Report structure")
    story = [Paragraph("QŪNAR", styles["title"]), Paragraph(combined_title, styles["subtitle"])]
    story.append(Paragraph(structure_title, styles["h2"]))
    for i, payload in enumerate(payloads, 1):
        story.append(Paragraph(f"{i}. {html.escape(payload.get('mode_label',''))}", styles["key"]))
    story.append(Spacer(1, 5*mm))
    for i, payload in enumerate(payloads):
        if i:
            story.append(PageBreak())
        _append_payload_pdf(story, payload, styles, include_inputs=False, include_brand=False)

    common, specifics = _consolidated_input_groups(payloads)
    if common or specifics:
        title = {"kk":"Біріктірілген бастапқы деректер мен дереккөздер","ru":"Сводные исходные данные и источники","en":"Consolidated inputs and sources"}.get(lang,"Consolidated inputs and sources")
        shared_title = {"kk":"Ортақ параметрлер","ru":"Общие параметры","en":"Shared parameters"}.get(lang,"Shared parameters")
        differences_title = {"kk":"Режимдер бойынша айырмашылықтар","ru":"Отличающиеся параметры по режимам","en":"Mode-specific inputs"}.get(lang,"Mode-specific inputs")
        story += [PageBreak(), Paragraph(title, styles['h2'])]
        def add_input_table(rows):
            labels = {'kk':('Көрсеткіш','Мәні','Дереккөз / рөл'),'ru':('Показатель','Значение','Источник / роль'),'en':('Parameter','Value','Source / role')}.get(lang,('Parameter','Value','Source / role'))
            data = [[Paragraph(html.escape(x),styles['small']) for x in labels]]
            data += [[Paragraph(html.escape(_clean(x)),styles['small']) for x in row] for row in rows]
            table = Table(data,colWidths=[52*mm,42*mm,76*mm],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF3EC')),('LINEBELOW',(0,0),(-1,-1),.25,colors.HexColor('#D7E0D9')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
            story.append(table)
        if common:
            story.append(Paragraph(shared_title,styles['h3']))
            add_input_table(common)
        if specifics:
            story.append(Paragraph(differences_title,styles['h3']))
            for label, rows in specifics:
                story.append(Paragraph(html.escape(label),styles['h3']))
                add_input_table(rows)
    doc.build(story)
    out.seek(0)
    return out.getvalue()


def _word_add_payload(doc, payload: dict, include_inputs: bool = True, include_brand: bool = True):
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def add_heading(text, level=1):
        p = doc.add_paragraph()
        p.style = doc.styles[f'Heading {min(level, 3)}']
        r = p.add_run(_clean(text))
        r.font.color.rgb = None
        return p

    if include_brand:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('QŪNAR')
        r.bold = True; r.font.size = Pt(24)
        if payload.get('brand_translation'):
            p2 = doc.add_paragraph(_clean(payload.get('brand_translation')))
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(payload.get('mode_label', ''), 1)
    meta = doc.add_paragraph()
    meta.add_run(f"{payload.get('generated_label','Generated')}: ").bold = True
    meta.add_run(_clean(payload.get('generated_at','')))
    if payload.get('assessment_year'):
        meta.add_run(f" | { {'kk':'Есептік жыл','ru':'Расчётный год','en':'Calculation year'}.get(payload.get('lang','en'),'Calculation year')}: {payload.get('assessment_year')}")
    if payload.get('assessment_type'):
        meta.add_run(f" | {_clean(payload.get('assessment_type'))}")

    summary = payload.get('summary', []) or []
    if summary:
        add_heading(payload.get('summary_title','Key results'), 2)
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, txt in enumerate([payload.get('indicator_label','Indicator'), payload.get('value_label','Value'), payload.get('note_label','Note')]):
            hdr[i].text = _clean(txt); shade(hdr[i], 'EAF3EC')
        for item in summary:
            cells = table.add_row().cells
            cells[0].text = _clean(item.get('label'))
            cells[1].text = _clean(item.get('value'))
            cells[2].text = _clean(item.get('note'))
            for c in cells: c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    recommendations = payload.get('recommendations', []) or []
    if recommendations:
        add_heading(payload.get('recommendations_title','Recommendations'), 2)
        for i, rec in enumerate(recommendations, 1):
            if isinstance(rec, dict):
                p = doc.add_paragraph(style='List Number')
                p.add_run(_clean(rec.get('title') or f'Recommendation {i}')).bold = True
                if rec.get('text') or rec.get('recommendation'):
                    doc.add_paragraph(_clean(rec.get('text') or rec.get('recommendation')))
                if rec.get('evidence'):
                    p = doc.add_paragraph(); p.add_run(f"{_clean(rec.get('evidence_label','Evidence'))}: ").bold = True; p.add_run(_clean(rec.get('evidence')))
                if rec.get('effect'):
                    p = doc.add_paragraph(); p.add_run(f"{_clean(rec.get('effect_label','Expected effect'))}: ").bold = True; p.add_run(_clean(rec.get('effect')))
            else:
                doc.add_paragraph(_clean(rec), style='List Number')

    analysis = payload.get('analysis', []) or []
    if analysis:
        add_heading(payload.get('analysis_title','Analysis'), 2)
        for text in analysis:
            doc.add_paragraph(_clean(text))

    table_data = payload.get('table')
    if table_data:
        add_heading(payload.get('table_title','Detailed result'), 2)
        df = pd.DataFrame(table_data)
        if not df.empty:
            df = df.iloc[:, :8].head(25)
            table = doc.add_table(rows=1, cols=len(df.columns))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, col in enumerate(df.columns):
                table.rows[0].cells[j].text = _clean(col); shade(table.rows[0].cells[j], 'EAF3EC')
            for _, rr in df.iterrows():
                cells = table.add_row().cells
                for j, val in enumerate(rr.tolist()): cells[j].text = _clean(val)

    if include_inputs and payload.get('inputs'):
        doc.add_page_break()
        add_heading(payload.get('inputs_title','Inputs and sources'), 2)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        labels = [payload.get('parameter_label','Parameter'), payload.get('value_label','Value'), payload.get('source_label','Source')]
        for j, txt in enumerate(labels): table.rows[0].cells[j].text = _clean(txt); shade(table.rows[0].cells[j], 'EAF3EC')
        for p_, v_, s_ in payload['inputs']:
            cells = table.add_row().cells
            cells[0].text = _clean(p_); cells[1].text = _clean(v_); cells[2].text = _clean(s_)


def _word_document_base(title='QŪNAR report'):
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8); section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)
    styles = doc.styles
    styles['Normal'].font.name = 'Arial'; styles['Normal'].font.size = Pt(11)
    styles['Normal'].paragraph_format.first_line_indent = Cm(0.65)
    styles['Normal'].paragraph_format.line_spacing = 1.35
    styles['Normal'].paragraph_format.space_after = Pt(7)
    for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        styles[name].font.name = 'Arial'; styles[name].font.size = Pt(size); styles[name].font.bold = True
        styles[name].paragraph_format.first_line_indent = Cm(0)
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.first_line_indent = Cm(0)
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.first_line_indent = Cm(0)
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.first_line_indent = Cm(0)
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(6)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for table_style_name in ('Table Grid',):
        if table_style_name in styles:
            styles[table_style_name].font.name = 'Arial'
            styles[table_style_name].font.size = Pt(9)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for table_style_name in ('Table Grid',):
        if table_style_name in styles:
            styles[table_style_name].font.name = 'Arial'
            styles[table_style_name].font.size = Pt(9)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for table_style_name in ('Table Grid',):
        if table_style_name in styles:
            styles[table_style_name].font.name = 'Arial'
            styles[table_style_name].font.size = Pt(9)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for table_style_name in ('Table Grid',):
        if table_style_name in styles:
            styles[table_style_name].font.name = 'Arial'
            styles[table_style_name].font.size = Pt(9)
    doc.core_properties.title = title
    doc.core_properties.author = 'QŪNAR'
    return doc


def build_word_report(payload: dict) -> bytes:
    doc = _word_document_base('QŪNAR report')
    _word_add_payload(doc, payload, include_inputs=True, include_brand=True)
    out = BytesIO(); doc.save(out); out.seek(0); return out.getvalue()


def build_combined_word_report(payloads: list[dict]) -> bytes:
    payloads = [p for p in payloads if p]
    lang = (payloads[0].get('lang') if payloads else 'en') or 'en'
    title = {'kk':'Біріктірілген аналитикалық есеп','ru':'Сводный аналитический отчет','en':'Integrated analytical report'}.get(lang,'Integrated analytical report')
    doc = _word_document_base('QŪNAR integrated report')
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('QŪNAR'); r.bold = True
    p = doc.add_paragraph(title); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    structure = {'kk':'Есеп құрылымы','ru':'Структура отчета','en':'Report structure'}.get(lang,'Report structure')
    doc.add_heading(structure, level=1)
    for i, payload in enumerate(payloads, 1): doc.add_paragraph(f"{i}. {_clean(payload.get('mode_label',''))}")
    for i, payload in enumerate(payloads):
        doc.add_page_break()
        _word_add_payload(doc, payload, include_inputs=False, include_brand=False)
    common, specifics = _consolidated_input_groups(payloads)
    if common or specifics:
        doc.add_page_break()
        doc.add_heading({'kk':'Біріктірілген бастапқы деректер мен дереккөздер','ru':'Сводные исходные данные и источники','en':'Consolidated inputs and sources'}.get(lang), level=1)
        def add_table(rows):
            table=doc.add_table(rows=1,cols=3); table.style='Table Grid'
            for j,x in enumerate(('Parameter','Value','Source / role')): table.rows[0].cells[j].text=x
            for p_,v_,s_ in rows:
                cells=table.add_row().cells;cells[0].text=_clean(p_);cells[1].text=_clean(v_);cells[2].text=_clean(s_)
        if common:
            doc.add_heading('Shared parameters' if lang=='en' else 'Общие параметры' if lang=='ru' else 'Ортақ параметрлер',level=2)
            add_table(common)
        if specifics:
            doc.add_heading('Mode-specific inputs' if lang=='en' else 'Отличающиеся параметры по режимам' if lang=='ru' else 'Режимдер бойынша айырмашылықтар',level=2)
            for label,rows in specifics:
                doc.add_heading(label,level=3);add_table(rows)
    out=BytesIO(); doc.save(out); out.seek(0); return out.getvalue()
