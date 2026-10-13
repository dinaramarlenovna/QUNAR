"""Cross-module regression: economic inputs, feasible crop plan and report integrity.

Only crop-allocation/margin/water-budget mechanics are in scope; no investment model.
"""
import pytest
from core.optimizer import optimize_structure
from core.optimization_guard import validate_allocation
from core.irrigation_economics import irrigation_comparison
from core.irrigation_report import irrigation_export
from core.reporting import build_pdf_report, build_excel_report, build_word_report

DISTRICT='Saryagash District'
MARGINS={
    'Картофель':300000,
    'Помидоры':500000,
    'Кукуруза (маис)':175000,
    'Пшеница':9000,
    'Хлопок, очищенный или не очищенный от семян':92700,
}
SHARES={
    'Картофель':15,
    'Помидоры':20,
    'Кукуруза (маис)':20,
    'Пшеница':25,
    'Хлопок, очищенный или не очищенный от семян':20,
}

def run(margins=None, shares=None, water=520000):
    return optimize_structure(DISTRICT,100,max_share=.25,n_crops=5,
        available_water_m3=water,economic_margin_by_product=MARGINS if margins is None else margins,
        objective_mode='balanced', baseline_share_by_product=SHARES if shares is None else shares)


def test_existing_reproducible_baseline_case_unchanged():
    alloc, out=run()
    assert out['water_feasible'] is True
    assert out['water_budget_verified'] is True
    assert out['baseline_status']=='valid_comparison'
    assert out['economic_gain_kzt']==pytest.approx(4274100,abs=1)
    assert out['baseline_contribution_margin_kzt']==pytest.approx(20079000,abs=1)
    assert out['optimized_contribution_margin_kzt']==pytest.approx(24353100,abs=1)
    assert alloc['modelled_water_m3'].sum()==pytest.approx(501700,abs=1)
    assert validate_allocation(alloc,out,area_ha=100,available_water_m3=520000,max_share=.25,n_crops=5)['ok']


def test_economic_priority_matches_existing_2027_pdf_scenario():
    alloc, out=optimize_structure(DISTRICT,100,profile='balanced',max_share=.25,n_crops=5,
        available_water_m3=520000,economic_margin_by_product=MARGINS,
        objective_mode='economic', economic_weight=.50,baseline_share_by_product=SHARES)
    assert out['optimized_contribution_margin_kzt']==pytest.approx(24851100,abs=1)
    assert out['economic_gain_kzt']==pytest.approx(4772100,abs=1)
    assert out['modelled_total_water_m3']==pytest.approx(508900,abs=1)


def test_baseline_partial_percentages_are_not_renormalized():
    alloc, info=run(shares={'Картофель':25,'Помидоры':20})
    assert not alloc.empty
    assert info['baseline_status']=='shares_do_not_sum_to_100'
    assert info['baseline_contribution_margin_kzt'] is None
    assert info['economic_gain_kzt'] is None


def test_baseline_over_100_is_not_renormalized():
    alloc, info=run(shares=dict(SHARES, **{'Картофель':60}))
    assert info['baseline_status']=='shares_do_not_sum_to_100'
    assert info['economic_gain_kzt'] is None


def test_normalized_fractions_still_supported():
    _, info=run(shares={k:v/100 for k,v in SHARES.items()})
    assert info['baseline_status']=='valid_comparison'
    assert info['economic_gain_kzt']==pytest.approx(4274100,abs=1)


def test_negative_margin_not_mistaken_for_missing():
    m=dict(MARGINS);m['Пшеница']=-9000
    alloc, info=run(margins=m)
    assert not alloc.empty
    assert info['economic_objective_active']
    assert info['economic_candidates_n']==5
    assert (alloc['economic_margin_kzt_ha']<0).any()
    assert info['baseline_status']=='valid_comparison'
    assert info['economic_gain_kzt'] is not None


def test_explicit_zero_margin_is_valid_break_even_observation():
    m=dict(MARGINS);m['Пшеница']=0
    alloc, info=run(margins=m)
    assert info['economic_candidates_n']==5
    assert info['baseline_status']=='valid_comparison'
    assert info['economic_gain_kzt'] is not None


def test_missing_margin_disables_comparison_without_fake_zero():
    m=dict(MARGINS);del m['Пшеница']
    alloc,info=run(margins=m)
    assert alloc.empty
    assert info['solver_status']=='infeasible_crop_count_or_cap'
    assert info.get('economic_gain_kzt') is None
    # Five crops requested but margins entered for only four. No imputation.
    assert len(m)==4
    assert info.get('economic_gain_kzt') is None


def test_no_water_input_never_proves_water_feasibility_or_margin_gain():
    alloc,info=run(water=0)
    assert not alloc.empty
    assert info['water_feasible'] is None
    assert info['water_budget_verified'] is False
    assert info['baseline_water_feasible'] is None
    assert info['baseline_status']=='water_budget_not_entered'
    assert info['economic_gain_kzt'] is None


def test_low_water_budget_rejects_unsupported_profit_claim():
    alloc,info=run(water=350000)
    assert alloc.empty
    assert info['water_feasible'] is False
    assert info.get('economic_gain_kzt') is None
    assert info.get('optimized_contribution_margin_kzt') is None


def test_realistic_zero_margin_and_large_loss_remain_finite():
    m=dict(MARGINS);m['Пшеница']=-3000000;m['Хлопок, очищенный или не очищенный от семян']=0
    alloc,info=run(margins=m)
    assert not alloc.empty
    assert info['optimized_contribution_margin_kzt'] is not None


def test_irrigation_report_risk_water_and_margin_values_agree_with_three_export_formats():
    comp=irrigation_comparison(
        baseline_yield_centners_ha=50,price_kzt_t=105000,
        other_variable_cost_kzt_ha=326000,baseline_water_m3_ha=6450,
        proposed_water_m3_ha=5160,baseline_energy_kwh_ha=420,
        proposed_energy_kwh_ha=420,water_tariff_kzt_m3=.845,
        electricity_tariff_kzt_kwh=45.9872,yield_change_pct=0,
        extra_technology_cost_kzt_ha=15000,evidence='QA scenario assumption')
    import fitz
    import io
    from docx import Document
    from openpyxl import load_workbook
    for lang in ('ru','kk','en'):
        rows, rec=irrigation_export(comp,area_ha=20,input_yield_centners_ha=50,
            ml_yield_centners_ha=63.94,selling_price_kzt_t=105000,
            other_variable_cost_kzt_ha=326000,extra_technology_cost_kzt_ha=15000,
            yield_change_pct=0,evidence='QA scenario assumption',lang=lang,
            available_water_m3=107000,reuse_pct=0)
        values={r['value'] for r in rows}
        assert '129,000' in values and '103,200' in values
        assert '-22,000' in values and '+3,800' in values
        payload={'mode_label':'Crop irrigation QA', 'summary_title':'Key results', 'summary':rows,
             'recommendations_title':'Recommendations','recommendations':[rec],
             'inputs_title':'Input assumptions', 'inputs':[['Year','2027','Scenario']]}
        pdf=build_pdf_report(payload)
        pdf_text='\n'.join(page.get_text() for page in fitz.open(stream=pdf,filetype='pdf'))
        assert '-22,000' in pdf_text and '+3,800' in pdf_text
        assert '129,000' in pdf_text and '103,200' in pdf_text
        xlsx=load_workbook(io.BytesIO(build_excel_report(payload)),read_only=True)
        excel_text=' '.join(str(c) for sh in xlsx for row in sh.values for c in row if c is not None)
        assert '-22,000' in excel_text and '+3,800' in excel_text
        docx=Document(io.BytesIO(build_word_report(payload)))
        word_text=' '.join(p.text for p in docx.paragraphs)+' '.join(cell.text for table in docx.tables for row in table.rows for cell in row.cells)
        assert '-22,000' in word_text and '+3,800' in word_text
