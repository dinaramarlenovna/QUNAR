from copy import deepcopy
from io import BytesIO
import json
import zipfile

from core.report_variants import (report_fingerprint, normalize_report_library,
    add_report, selected_reports, compare_pair, compared_variants, condensed_variants)
from core.reporting import (build_combined_pdf_report, build_combined_word_report,
    build_combined_excel_report)


def example(used=35500, water=45000):
    return {'mode_key':'mode1', 'lang':'ru','generated_at':'2026-10-18 16:00',
        'mode_label':'Оценить моё хозяйство', 'assessment_year':2027,
        'inputs':[['Площадь культуры, га','10','Ввод пользователя'],
                  ['Доступная вода, м³',str(water),'Ввод пользователя']],
        'summary':[{'label':'Потребность воды, м³','value':str(used)}],
        'recommendations':[{'title':'Водная обеспеченность','text':'Проверить лимит воды'}],
        'table':[{'Культура':'Пшеница','Вода':used}],
        'analysis':['Сценарные данные, не фактические измерения.']}


def test_duplicate_ignores_generation_time():
    a=example()
    b=deepcopy(a);b['generated_at']='2026-10-18 16:44'
    basket={}
    assert add_report(basket,a)
    assert not add_report(basket,b)
    assert len(basket)==1
    assert report_fingerprint(a)==report_fingerprint(b)


def test_distinct_scenarios_same_mode_are_preserved():
    basket={}
    assert add_report(basket,example())
    assert add_report(basket,example(used=28400,water=40000))
    assert len(basket)==2
    pairs=compared_variants(list(basket.values()))
    assert len(pairs)==1
    assert ('Потребность воды, м³','35500','28400') in pairs[0]['results']
    assert ('Доступная вода, м³','45000','40000') in pairs[0]['inputs']


def test_checkbox_excludes_without_deleting():
    basket={}
    add_report(basket,example())
    add_report(basket,example(28400))
    ids=list(basket)
    select=selected_reports(basket,{ids[0]:False,ids[1]:True})
    assert len(select)==1 and len(basket)==2
    assert select[0]['summary'][0]['value']=='28400'
    basket.pop(ids[0]);assert len(basket)==1


def test_legacy_basket_migration_and_condensing():
    legacy={'mode1':example()}
    recovered=normalize_report_library(legacy)
    assert len(recovered)==1 and 'mode1' not in recovered
    source=[example(),example(used=28400,water=40000)]
    saved=condensed_variants(source)
    assert len(saved[0]['inputs'])==2
    assert len(saved[1]['inputs'])==1
    assert saved[1]['table']!=saved[0]['table']
    assert len(saved[1]['summary'])==1


def test_actual_exports_contain_comparison_and_units():
    samples=[example(),example(28400,40000)]
    pdf=build_combined_pdf_report(samples)
    assert pdf.startswith(b'%PDF') and len(pdf)>1000
    docx=build_combined_word_report(samples)
    assert zipfile.is_zipfile(BytesIO(docx))
    from docx import Document
    doc=Document(BytesIO(docx))
    all_text=' '.join(p.text for p in doc.paragraphs)
    assert 'Сравнение вариантов' in all_text
    assert 'Потребность воды' in ' '.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    xlsx=build_combined_excel_report(samples)
    assert zipfile.is_zipfile(BytesIO(xlsx))
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(xlsx),read_only=True)
    assert 'Comparisons' in wb.sheetnames
    assert len(wb.sheetnames)>=4
