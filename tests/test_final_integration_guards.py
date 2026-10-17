"""Critical scientific claims and new scenario unit regression tests."""
import pytest
from pathlib import Path
from core.cmip6_scenarios import load_scenario
from core.carbon_scenarios import partial_inventory, compare_partial_inventory
from core.nasa_polygon import polygon_sample_locations, fetch_polygon_growing_season
from core.yield_reference import trailing_three_year_yield
from core.optimizer import optimize_structure
from core.optimization_guard import validate_allocation
from shapely.geometry import Polygon

DATA = Path(__file__).resolve().parents[1] / 'data'


def test_cmip6_scenarios_from_uploaded_csv_do_not_mix_mean_and_anomalies():
    for s, temp, rainfall_mm, days35, days40, days45 in [
        ('SSP2-4.5',2.01,1.59,63.04,16.55,.54),
        ('SSP5-8.5',2.68,1.03,70.41,21.45,1.01),
    ]:
        x=load_scenario(DATA/'cmip6',s)
        assert x.temperature_change_c==pytest.approx(temp,abs=.025)
        assert x.precipitation_change_mm==pytest.approx(rainfall_mm,abs=.025)
        assert x.precipitation_change_pct==pytest.approx(100*rainfall_mm/121.47,abs=.03)
        assert sum(x.days_over_35_monthly)==pytest.approx(days35,abs=.005)
        assert sum(x.days_over_40_monthly)==pytest.approx(days40,abs=.005)
        assert sum(x.days_over_45_monthly)==pytest.approx(days45,abs=.005)


def test_partial_inventory_no_unsupported_enteric_zero_when_heads_present():
    with pytest.raises(ValueError,match='factor must be positive'):
        partial_inventory(nitrogen_kg=10,diesel_litres=10,electricity_kwh=10,
                          grid_kg_co2e_kwh=.8,include_enteric=True,
                          enteric_heads=5,enteric_kg_ch4_head_year=0,enteric_year_fraction=1)


def test_comparisons_require_same_boundary_and_complete_activities():
    a=partial_inventory(nitrogen_kg=10,diesel_litres=10,electricity_kwh=10,grid_kg_co2e_kwh=.8)
    b=partial_inventory(nitrogen_kg=10,diesel_litres=10,electricity_kwh=10,grid_kg_co2e_kwh=.8,
                        include_enteric=True,enteric_heads=1,enteric_kg_ch4_head_year=50,enteric_year_fraction=1)
    with pytest.raises(ValueError,match='Incompatible'):
        compare_partial_inventory(a,b)
    with pytest.raises(ValueError,match='Matched complete'):
        compare_partial_inventory(a,partial_inventory(nitrogen_kg=None,diesel_litres=10,electricity_kwh=10,grid_kg_co2e_kwh=.8))


def test_user_polygon_same_grid_one_nasa_point():
    poly=Polygon([(68.30,42.10),(68.31,42.10),(68.31,42.11),(68.30,42.11)])
    points=polygon_sample_locations(poly)
    assert len(points)==1 and points[0].area_weight==pytest.approx(1)


def test_user_polygon_no_partial_api_average_if_one_cell_fails():
    poly=Polygon([(68.4,42.25),(69.0,42.25),(69.0,42.6),(68.4,42.6)])
    calls=0
    def fake(lat,lon,year):
        nonlocal calls
        calls+=1
        if calls==2:
            return {'Apr-Sep temperature, °C':15}
        return {
            'Apr-Sep temperature, °C':20,
            'Apr-Sep precipitation, mm':100,
            'Apr-Sep relative humidity, %':40,
            'Apr-Sep root-zone wetness':.2,
            'Apr-Sep solar radiation, MJ/m²/day':15,
            'period_start':'2024-04-01','period_end':'2024-09-30',
        }
    with pytest.raises(ValueError,match='Incomplete'):
        fetch_polygon_growing_season(poly,2024,fetch_point=fake)
    assert calls==2


def test_custom_weight_profile_has_no_silent_relaxation():
    d, info=optimize_structure('Saryagash District',100,profile='custom',custom_climate_weight=.15,
                     max_share=.25,n_crops=5,available_water_m3=520000)
    assert not d.empty
    assert info['climate_weight']==pytest.approx(.15)
    assert info['water_weight']==pytest.approx(.85)
    assert info['profile_source']=='user sensitivity, not survey-calibrated'
    assert validate_allocation(d,info,area_ha=100,available_water_m3=520000,max_share=.25,n_crops=5)['ok']
    d2, info2=optimize_structure('Saryagash District',100,profile='custom',custom_climate_weight=.15,
                        max_share=.25,n_crops=5,available_water_m3=250000)
    assert d2.empty
    assert info2['water_feasible'] is False


def test_future_yield_reference_never_recycles_2025_for_2025_forecast():
    r=trailing_three_year_yield(DATA/'analytical_panel.csv','Saryagash District','Wheat',2025)
    assert all(y < 2025 for y in r.years)
    assert set(r.requested_years)=={2022,2023,2024}

@pytest.mark.parametrize('language,caption',[
    ('ru','Изменение климата'),
    ('kk','Климаттың өзгеруі'),
    ('en','Climate change'),
])
def test_scenario_source_survives_all_three_report_exports(language, caption):
    # Test actual PDF/DOCX/XLSX engines using a representative payload.
    from io import BytesIO
    from docx import Document
    import fitz
    from core.reporting import build_pdf_report, build_word_report, build_excel_report
    payload = {
        'lang': language, 'mode_label': caption,
        'summary': [
            {'label': caption, 'value': '+2.01 °C', 'note': 'CMIP6 SSP2-4.5, 2040-2059'},
            {'label': 'Partial GHG, kg CO2e', 'value': '400.00', 'note': 'Model scenario, not measured reduction'},
        ],
        'inputs': [
            ['CMIP6 scenario','SSP2-4.5 / 2040-2059','World Bank CCKP, 1995-2014 baseline'],
            ['GHG inputs','N 100 kg; diesel 80 L; electricity 500 kWh','User-declared example'],
        ],
    }
    pdf=build_pdf_report(payload)
    assert pdf.startswith(b'%PDF')
    with fitz.open(stream=pdf,filetype='pdf') as d:
        txt=' '.join(p.get_text() for p in d)
        assert 'SSP2-4.5' in txt and '400.00' in txt
        assert '1995-2014' in txt
    word=Document(BytesIO(build_word_report(payload)))
    wordtxt=' '.join([p.text for p in word.paragraphs]+[cell.text for t in word.tables for row in t.rows for cell in row.cells])
    assert 'SSP2-4.5' in wordtxt and '400.00' in wordtxt and '1995-2014' in wordtxt
    xls=build_excel_report(payload)
    import zipfile
    with zipfile.ZipFile(BytesIO(xls)) as z:
        assert z.testzip() is None
        shared=''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.endswith('sharedStrings.xml'))
        sheets=''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
        assert 'SSP2-4.5' in shared+sheets and '1995-2014' in shared+sheets
